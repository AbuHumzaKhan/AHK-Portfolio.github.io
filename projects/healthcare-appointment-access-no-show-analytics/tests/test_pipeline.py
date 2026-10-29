from pathlib import Path
import json, sqlite3, unittest
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]

class PipelineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.df = pd.read_csv(ROOT / "data" / "appointments_synthetic.csv")
        cls.metrics = json.loads((ROOT / "outputs" / "metrics_summary.json").read_text(encoding="utf-8"))

    def test_expected_record_count_and_date_window(self):
        self.assertEqual(len(self.df), 7200)
        self.assertEqual(self.df.appointment_date.min(), "2026-04-01")
        self.assertEqual(self.df.appointment_date.max(), "2026-09-27")

    def test_primary_key_unique(self):
        self.assertTrue(self.df.appointment_id.is_unique)

    def test_booking_not_after_appointment(self):
        self.assertTrue((pd.to_datetime(self.df.appointment_date) >= pd.to_datetime(self.df.booking_date)).all())

    def test_status_domain(self):
        self.assertEqual(set(self.df.status.unique()), {"Attended", "No-show", "Cancelled", "Rescheduled"})

    def test_wait_time_rule(self):
        self.assertTrue(self.df.loc[self.df.status.eq("Attended"), "wait_time_min"].notna().all())
        self.assertTrue(self.df.loc[~self.df.status.eq("Attended"), "wait_time_min"].isna().all())

    def test_no_show_denominator_definition(self):
        eligible = self.df.status.isin(["Attended", "No-show"]).sum()
        no_show = self.df.status.eq("No-show").sum()
        self.assertAlmostEqual(self.metrics["kpis"]["no_show_rate_pct"], round(no_show / eligible * 100, 2))

    def test_sqlite_integrity(self):
        conn = sqlite3.connect(ROOT / "data" / "appointments.sqlite")
        self.assertEqual(conn.execute("PRAGMA integrity_check").fetchone()[0], "ok")
        self.assertEqual(conn.execute("SELECT COUNT(*) FROM appointments").fetchone()[0], 7200)
        conn.close()

    def test_qa_all_passed(self):
        qa = pd.read_csv(ROOT / "outputs" / "qa_checks.csv")
        self.assertTrue(qa.passed.all())

    def test_excel_has_formulas_and_charts(self):
        wb = load_workbook(ROOT / "excel" / "Appointment_Access_NoShow_Analysis.xlsx", data_only=False)
        self.assertIn("=COUNTA(Appointments!$A:$A)-1", [wb["Executive Summary"].cell(r,2).value for r in range(4,12)])
        self.assertGreaterEqual(len(wb["Clinic Summary"]._charts), 1)
        self.assertGreaterEqual(len(wb["Specialty Summary"]._charts), 1)
        self.assertGreaterEqual(len(wb["Lead Time Summary"]._charts), 1)

    def test_dashboard_assets_exist(self):
        html = (ROOT / "reports" / "dashboard.html").read_text(encoding="utf-8")
        self.assertIn("Synthetic data notice", html)
        self.assertNotIn("DATA_JSON", html)
        self.assertIn("Appointment Access & No-show Analytics", html)

    def test_seed_sql_exists_and_has_inserts(self):
        seed = (ROOT / "sql" / "seed_data.sql").read_text(encoding="utf-8")
        self.assertIn("INSERT INTO clinics", seed)
        self.assertGreaterEqual(seed.count("INSERT INTO appointments"), 120)

if __name__ == "__main__": unittest.main(verbosity=2)
